"""Step 1 문서 수집·정규화 — PDF / HWPX / DOCX / TXT 에서 텍스트·표를 뽑고 민감정보를 가린다.

- LLM 에는 원본 파일이 아니라 마스킹을 거친 추출 텍스트만 보낸다 (개인정보 반출 통제).
- 표는 셀 구조를 ' | ' 로 유지한다 (요금·수치·조항이 표 안에 있는 경우가 많다).
- HWP(구 바이너리)는 hwp5txt 가 설치돼 있을 때만 읽는다. 없으면 HWPX/PDF 로 저장해 올리라고 안내한다.
- 스캔본 OCR 은 범위 밖. 텍스트가 거의 없으면 경고를 남긴다.
"""
from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Extracted:
    text: str
    pages: int = 0
    warnings: list[str] = field(default_factory=list)
    masked: dict = field(default_factory=dict)


class IngestError(Exception):
    pass


def _pdf(data: bytes) -> Extracted:
    import pdfplumber

    parts, warnings = [], []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        n = len(pdf.pages)
        for i, page in enumerate(pdf.pages, 1):
            body = page.extract_text() or ""
            tables = []
            for t in page.extract_tables() or []:
                rows = [" | ".join((c or "").replace("\n", " ").strip() for c in r) for r in t if r]
                if rows:
                    tables.append("\n".join(rows))
            chunk = f"[p.{i}]\n{body}"
            if tables:
                chunk += "\n[표]\n" + "\n\n".join(tables)
            parts.append(chunk)
    text = "\n\n".join(parts)
    if len(re.sub(r"\s|\[p\.\d+\]", "", text)) < 50 * max(n, 1) / 4:
        warnings.append("텍스트 레이어가 거의 없습니다. 스캔본이면 OCR 후 다시 올려 주세요.")
    return Extracted(text=text, pages=n, warnings=warnings)


def _hwpx(data: bytes) -> Extracted:
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise IngestError("HWPX 파일을 열 수 없습니다.") from e
    sections = sorted(n for n in z.namelist() if re.match(r"Contents/section\d+\.xml", n))
    paras = []
    for s in sections:
        xml = z.read(s).decode("utf-8", "ignore")
        for p in re.findall(r"<hp:p\b.*?</hp:p>", xml, re.S):
            t = "".join(re.findall(r"<hp:t[^>]*>(.*?)</hp:t>", p, re.S))
            t = re.sub(r"<[^>]+>", "", t)
            if t.strip():
                paras.append(_unescape(t.strip()))
    return Extracted(text="\n".join(paras), pages=len(sections))


def _hwp(data: bytes) -> Extracted:
    exe = shutil.which("hwp5txt")
    if not exe:
        raise IngestError("HWP(구 형식)는 hwp5txt(pyhwp)가 있어야 읽을 수 있습니다. "
                          "한글에서 HWPX 또는 PDF 로 저장해 올려 주세요.")
    with tempfile.NamedTemporaryFile(suffix=".hwp", delete=False) as f:
        f.write(data)
        tmp = f.name
    try:
        out = subprocess.run([exe, tmp], capture_output=True, timeout=120)
    finally:
        Path(tmp).unlink(missing_ok=True)
    if out.returncode != 0:
        raise IngestError("HWP 변환 실패: " + out.stderr.decode("utf-8", "ignore")[:200])
    return Extracted(text=out.stdout.decode("utf-8", "ignore"))


def _docx(data: bytes) -> Extracted:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        parts.append("[표]\n" + "\n".join(" | ".join(c.text.strip() for c in r.cells) for r in t.rows))
    return Extracted(text="\n".join(parts))


def _unescape(s: str) -> str:
    return (s.replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
             .replace("&apos;", "'").replace("&amp;", "&"))


def _html(data: bytes) -> Extracted:
    s = data.decode("utf-8", "ignore")
    s = re.sub(r"(?is)<(script|style).*?</\1>", "", s)
    s = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>|</tr>", "\n", s)
    s = re.sub(r"<[^>]+>", "", s)
    return Extracted(text=_unescape(re.sub(r"\n\s*\n+", "\n\n", s)).strip())


def extract(filename: str, data: bytes) -> Extracted:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        ex = _pdf(data)
    elif ext == ".hwpx":
        ex = _hwpx(data)
    elif ext == ".hwp":
        ex = _hwp(data)
    elif ext == ".docx":
        ex = _docx(data)
    elif ext in (".html", ".htm"):
        ex = _html(data)
    elif ext in (".txt", ".md", ""):
        ex = Extracted(text=data.decode("utf-8", "ignore"))
    else:
        raise IngestError(f"지원하지 않는 형식: {ext} (PDF, HWPX, HWP, DOCX, HTML, TXT)")
    ex.text = normalize(ex.text)
    if not ex.text.strip():
        raise IngestError("문서에서 텍스트를 추출하지 못했습니다.")
    return ex


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ------------------------------------------------------------- 마스킹 규칙
# 워크북 샘플의 「OOO」 표기 관행을 시스템 규칙으로 승격 (XI장 3절)
_MASKS = [
    ("주민등록번호", re.compile(r"\b\d{6}-?[1-4]\d{6}\b"), "******-*******"),
    ("전화번호", re.compile(r"\b01[016789][-. ]?\d{3,4}[-. ]?\d{4}\b"), "010-****-****"),
    ("전화번호", re.compile(r"\b0\d{1,2}[-. ]\d{3,4}[-. ]\d{4}\b"), "0**-****-****"),
    ("이메일", re.compile(r"\b[\w.+-]+@[\w-]+\.[\w.-]+\b"), "***@***"),
    ("계좌번호", re.compile(r"\b\d{3,6}-\d{2,6}-\d{3,8}\b"), "***-***-****"),
]


def mask_pii(text: str) -> tuple[str, dict]:
    counts: dict[str, int] = {}
    for name, rx, repl in _MASKS:
        text, n = rx.subn(repl, text)
        if n:
            counts[name] = counts.get(name, 0) + n
    return text, counts
