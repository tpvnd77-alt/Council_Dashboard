"""Step 5 Executive Delivery — 10단계 JSON 을 Decision Log 텍스트 컬럼, 원페이지 보고서(HTML/DOCX)로 변환."""
from __future__ import annotations

import html
import io

from .config import Taxonomy, codebook
from .scoring import AREAS, d_day
from .schema import DecisionReport

AREA_ORDER = ["SCORE_REG", "SCORE_FIN", "SCORE_PR", "SCORE_LEG"]


def decision_log_fields(r: DecisionReport, panel: dict, tx: Taxonomy) -> dict:
    """IX장 매핑표대로 10단계 출력을 Decision Log 텍스트 컬럼에 나눠 담는다."""
    facts = "\n".join(f"· {f.fact}" for f in r.key_content.key_facts[:3])
    risk = [f"[표면] {r.real_issue.surface_issue}", f"[실제 쟁점] {r.real_issue.real_issue}",
            f"[상대방 의도] {r.real_issue.counterpart_intent}", f"[놓치면 안 되는 것] {r.real_issue.must_not_miss}"]
    for k in AREA_ORDER:
        if panel and k in panel:
            risk.append(f"[{AREAS[k][0]} {panel[k]['score']}] {panel[k]['summary']}")
    rec = next((o for o in r.decision.options if o.key == r.decision.recommended), None)
    dp = [f"Option {o.key}. {o.title} — {o.description}" for o in r.decision.options]
    dp.append(f"★ 추천 {r.decision.recommended}: {r.decision.recommended_summary}")
    dp.append(f"왜 지금: {r.decision.why_now} / 얻는 것: {r.decision.gain} / 감수: {r.decision.cost_to_accept}")
    return {
        "EXECUTIVE_SUMMARY": r.one_line_judgment + ("\n" + facts if facts else ""),
        "RISK_DESCRIPTION": "\n".join(risk),
        "DECISION_POINTS": "\n".join(dp),
        "TODO_ACTIONS": "\n".join(f"{a.when} [{a.owner}] {a.task} → {a.counterpart}: {a.message}"
                                  + (f" (결정: {a.decision_needed})" if a.decision_needed else "")
                                  for a in r.action_plan),
        "PR_STANCE": f"공식: 「{r.pr_stance.official_message}」 / 강조: {', '.join(r.pr_stance.emphasize)}"
                     f" / 금기어: {', '.join(r.pr_stance.taboo_words)}",
        "SCENARIOS": "\n".join(f"{s.name}: 효과 {s.effect} / 부작용 {s.side_effect}" for s in r.scenarios),
        "NEGOTIATION_CARDS": "\n".join([
            "지킬 것: " + "; ".join(r.negotiation_cards.protect),
            "양보 가능: " + "; ".join(r.negotiation_cards.concede),
            "요구할 것: " + "; ".join(r.negotiation_cards.demand),
            "Big Deal: " + "; ".join(r.negotiation_cards.big_deal)]),
        "CEO_QA": "\n".join(f"Q{i+1}. {q.question}\nA. {q.answer}" for i, q in enumerate(r.ceo_qa)),
        "AGENT_RECOMMENDATION": f"★ {r.decision.recommended}. {rec.title if rec else ''} — {r.decision.recommended_summary}",
    }


# ------------------------------------------------------------- HTML 원페이지
def _e(s) -> str:
    return html.escape(str(s if s is not None else ""))


def _li(items) -> str:
    return "".join(f"<li>{_e(x)}</li>" for x in items)


def report_html(rec: dict, standalone: bool = True) -> str:
    r = DecisionReport.model_validate(rec["REPORT_JSON"])
    panel = rec.get("PANEL_JSON") or {}
    tx = Taxonomy()
    cb = codebook()
    lvl = next((x for x in cb["levels"] if x["level"] == rec.get("RISK_LEVEL")), {})
    dd = d_day(rec.get("DUE_DATE"))
    rows = []
    for k in AREA_ORDER + ["SCORE_OPP"]:
        p = panel.get(k, {})
        name = AREAS[k][0] if k in AREAS else "전략적 기회"
        ev = "<br>".join(f"<span class='tag {'est' if e['kind']=='추정' else 'src'}'>{_e(e['kind'])}</span> "
                         f"{_e(e['quote'] or e['reasoning'])}" for e in p.get("evidence", [])[:2])
        rows.append(f"<tr><th>{name}</th><td class='sc'>{_e(rec.get(k, p.get('score')))}</td>"
                    f"<td>{_e(p.get('summary'))}<div class='ev'>{ev}</div></td></tr>")
    opts = "".join(
        f"<div class='opt {'rec' if o.key == r.decision.recommended else ''}'>"
        f"<b>{'★ ' if o.key == r.decision.recommended else ''}Option {o.key}. {_e(o.title)}</b><p>{_e(o.description)}</p>"
        f"<div class='pc'><div>＋ {'<br>＋ '.join(map(_e, o.pros))}</div><div>－ {'<br>－ '.join(map(_e, o.cons))}</div></div></div>"
        for o in sorted(r.decision.options, key=lambda o: o.key))
    ra = r.decision.rationale
    actions = "".join(f"<tr><td><b>{a.when}</b></td><td>{_e(a.owner)}</td><td>{_e(a.task)}</td>"
                      f"<td>{_e(a.counterpart)}</td><td>{_e(a.message)}</td><td>{_e(a.decision_needed)}</td></tr>"
                      for a in r.action_plan)
    body = f"""
<article class="onepager">
<header>
  <div class="hd1"><span class="lv lv-{_e(rec.get('RISK_LEVEL'))}">{_e(rec.get('RISK_LEVEL'))} · {_e(lvl.get('grade',''))}</span>
  <b>{_e(rec.get('DOC_TITLE'))}</b></div>
  <div class="hd2">{_e(rec['DOC_ID'])} · {_e(rec.get('SOURCE_TYPE'))} · {_e(tx.label(rec.get('ISSUE_ID') or ''))}
  · 종합 {_e(rec.get('TOTAL_RISK_SCORE'))}점 · 기회 {_e(rec.get('SCORE_OPP'))} · 보고 Level <b>{_e(rec.get('DECISION_TYPE'))}</b>
  · 1차 기한 {_e(rec.get('DUE_DATE'))}{f' (D{dd:+d})' if dd is not None else ''} · {_e(rec.get('STATUS'))}</div>
</header>
<section class="judg">① {_e(r.one_line_judgment)}</section>
<section><h3>② 핵심 내용</h3><ul>{''.join(f"<li>{_e(f.fact)}<div class='q'>「{_e(f.quote)}」</div></li>" for f in r.key_content.key_facts)}</ul>
<p><b>핵심 변화</b> {_e(r.key_content.core_change)}</p><p><b>회사 영향</b> {' › '.join(map(_e, r.key_content.company_impact))}</p></section>
<section><h3>③ 조직장이 봐야 할 진짜 쟁점</h3><table class="kv">
<tr><th>표면적 이슈</th><td>{_e(r.real_issue.surface_issue)}</td></tr><tr><th>실제 쟁점</th><td>{_e(r.real_issue.real_issue)}</td></tr>
<tr><th>상대방 의도</th><td>{_e(r.real_issue.counterpart_intent)}</td></tr><tr><th>놓치면 안 되는 것</th><td>{_e(r.real_issue.must_not_miss)}</td></tr></table></section>
<section><h3>④ Risk / Opportunity</h3><table class="risk">{''.join(rows)}</table><p class="note">기회: {_e(r.opportunity_note)}</p></section>
<section><h3>⑤ 향후 Scenario</h3><table class="kv">{''.join(f"<tr><th>{_e(s.name)}</th><td>{_e(s.effect)}<div class='ev'>부작용: {_e(s.side_effect)}</div></td></tr>" for s in r.scenarios)}</table></section>
<section><h3>⑥ 조직장 의사결정</h3><div class="opts">{opts}</div>
<p class="recline">★ 추천 {r.decision.recommended} — {_e(r.decision.recommended_summary)}</p>
<p><b>왜 지금</b> {_e(r.decision.why_now)} · <b>얻는 것</b> {_e(r.decision.gain)} · <b>감수할 것</b> {_e(r.decision.cost_to_accept)}</p>
<table class="kv small"><tr><th>손익</th><td>{_e(ra.pnl)}</td></tr><tr><th>정책</th><td>{_e(ra.policy)}</td></tr><tr><th>여론</th><td>{_e(ra.public_opinion)}</td></tr>
<tr><th>법무</th><td>{_e(ra.legal)}</td></tr><tr><th>전략</th><td>{_e(ra.strategy)}</td></tr></table></section>
<section><h3>⑦ 협상 카드</h3><div class="cards"><div><b>지킬 것</b><ul>{_li(r.negotiation_cards.protect)}</ul></div>
<div><b>양보 가능</b><ul>{_li(r.negotiation_cards.concede)}</ul></div><div><b>요구할 것</b><ul>{_li(r.negotiation_cards.demand)}</ul></div>
<div><b>Big Deal</b><ul>{_li(r.negotiation_cards.big_deal)}</ul></div></div></section>
<section><h3>⑧ Action Plan</h3><table class="act"><tr><th></th><th>담당</th><th>할 일</th><th>만날 사람</th><th>메시지</th><th>필요 결정</th></tr>{actions}</table></section>
<section class="att"><h3>[첨부 1] ⑨ PR Stance</h3><p><b>공식</b> 「{_e(r.pr_stance.official_message)}」</p>
<p><b>강조</b> {' / '.join(map(_e, r.pr_stance.emphasize))}</p><p><b>금기어</b> <span class="taboo">{' · '.join(map(_e, r.pr_stance.taboo_words))}</span></p></section>
<section class="att"><h3>[첨부 2] ⑩ CEO 예상 질문</h3><ol class="qa">{''.join(f"<li><b>{_e(q.question)}</b><br>{_e(q.answer)}</li>" for q in r.ceo_qa)}</ol></section>
{f"<section class='att'><h3>추정·확인 필요</h3><ul>{_li(r.assumptions + r.missing_numbers)}</ul></section>" if (r.assumptions or r.missing_numbers) else ''}
<footer>PROMPT {_e(rec.get('PROMPT_VERSION'))} · MODEL {_e(rec.get('MODEL_VERSION'))} · 검수 {_e(rec.get('REVIEWER') or '-')} · 결정 {_e(rec.get('APPROVER') or '-')}
{f"<br>조직장 결정: {_e(rec.get('FINAL_DECISION'))}" if rec.get('FINAL_DECISION') else ''}</footer>
</article>"""
    if not standalone:
        return body
    return f"""<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>{_e(rec['DOC_ID'])} 의사결정 보고서</title>
<style>{REPORT_CSS}</style></head><body>{body}</body></html>"""


REPORT_CSS = """
.onepager{font-family:'Noto Sans KR',system-ui,sans-serif;font-size:12.5px;line-height:1.5;color:#1b1f2a;max-width:860px;margin:0 auto;padding:16px}
.onepager header{border-bottom:2px solid #1b1f2a;padding-bottom:6px;margin-bottom:8px}
.onepager .hd1{font-size:16px;display:flex;gap:8px;align-items:center}.onepager .hd2{color:#555;font-size:11.5px;margin-top:2px}
.onepager .lv{font-size:11px;padding:2px 8px;border-radius:10px;color:#fff;background:#888;white-space:nowrap}
.onepager .lv-CRITICAL{background:#c62828}.onepager .lv-HIGH{background:#ef6c00}.onepager .lv-MEDIUM{background:#b59a00}.onepager .lv-LOW{background:#2e7d32}
.onepager .judg{font-size:14px;font-weight:700;background:#f3f5fb;border-left:4px solid #3949ab;padding:8px 10px;margin:8px 0}
.onepager h3{font-size:13px;margin:12px 0 4px;color:#283593}.onepager section{break-inside:avoid}
.onepager ul,.onepager ol{margin:2px 0;padding-left:18px}.onepager p{margin:3px 0}
.onepager .q,.onepager .ev{color:#666;font-size:11px}
.onepager table{border-collapse:collapse;width:100%}.onepager td,.onepager th{border:1px solid #ddd;padding:3px 6px;vertical-align:top;text-align:left}
.onepager .kv th,.onepager .risk th{width:110px;background:#fafafa;white-space:nowrap}.onepager .sc{width:28px;text-align:center;font-weight:700;font-size:14px}
.onepager .tag{font-size:10px;padding:0 4px;border-radius:3px;border:1px solid #999}.onepager .tag.est{color:#a15c00;border-color:#e0a050}
.onepager .opts{display:grid;grid-template-columns:1fr 1fr;gap:8px}.onepager .opt{border:1px solid #ccc;border-radius:6px;padding:6px 8px}
.onepager .opt.rec{border:2px solid #3949ab;background:#f6f7ff}.onepager .pc{display:grid;grid-template-columns:1fr 1fr;gap:6px;font-size:11.5px}
.onepager .recline{font-weight:700;color:#283593;margin-top:6px}.onepager .small{font-size:11.5px}
.onepager .cards{display:grid;grid-template-columns:repeat(4,1fr);gap:6px}.onepager .cards>div{background:#fafafa;border-radius:6px;padding:4px 6px}
.onepager .act th{background:#fafafa}.onepager .taboo{color:#c62828}.onepager .note{color:#555}
.onepager .att{border-top:1px dashed #aaa;margin-top:10px}.onepager footer{margin-top:12px;color:#888;font-size:10.5px;border-top:1px solid #ddd;padding-top:4px}
@media (max-width:640px){.onepager .opts,.onepager .cards{grid-template-columns:1fr}}
@media print{.onepager{padding:0}}
"""


# ------------------------------------------------------------------ DOCX
def report_docx(rec: dict) -> bytes:
    import docx
    from docx.shared import Pt

    r = DecisionReport.model_validate(rec["REPORT_JSON"])
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.size = Pt(9.5)
    st.font.name = "맑은 고딕"
    d.add_heading(rec.get("DOC_TITLE") or rec["DOC_ID"], level=1)
    d.add_paragraph(f"{rec['DOC_ID']} · {rec.get('SOURCE_TYPE')} · {Taxonomy().label(rec.get('ISSUE_ID') or '')} · "
                    f"{rec.get('RISK_LEVEL')} {rec.get('TOTAL_RISK_SCORE')}점 · 보고 Level {rec.get('DECISION_TYPE')} · "
                    f"기한 {rec.get('DUE_DATE')}")
    d.add_paragraph("① " + r.one_line_judgment).runs[0].bold = True
    sections = [
        ("② 핵심 내용", [f.fact for f in r.key_content.key_facts] + ["핵심 변화: " + r.key_content.core_change]),
        ("③ 진짜 쟁점", [f"표면: {r.real_issue.surface_issue}", f"실제: {r.real_issue.real_issue}",
                       f"상대방 의도: {r.real_issue.counterpart_intent}", f"놓치면 안 되는 것: {r.real_issue.must_not_miss}"]),
        ("④ Risk / Opportunity", [f"{AREAS[k][0]} {rec.get(k)}" for k in AREA_ORDER] +
         [f"전략적 기회 {rec.get('SCORE_OPP')} — {r.opportunity_note}"]),
        ("⑤ Scenario", [f"{s.name}: {s.effect} (부작용: {s.side_effect})" for s in r.scenarios]),
        ("⑥ 조직장 의사결정", [f"Option {o.key}. {o.title} — {o.description}" for o in r.decision.options] +
         [f"★ 추천 {r.decision.recommended}: {r.decision.recommended_summary}", f"왜 지금: {r.decision.why_now}"]),
        ("⑦ 협상 카드", ["지킬 것: " + "; ".join(r.negotiation_cards.protect), "양보: " + "; ".join(r.negotiation_cards.concede),
                       "요구: " + "; ".join(r.negotiation_cards.demand), "Big Deal: " + "; ".join(r.negotiation_cards.big_deal)]),
    ]
    for h, items in sections:
        d.add_heading(h, level=2)
        for it in items:
            d.add_paragraph(it, style="List Bullet")
    d.add_heading("⑧ Action Plan", level=2)
    t = d.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for i, h in enumerate(["시점", "담당", "할 일", "만날 사람", "메시지"]):
        t.rows[0].cells[i].text = h
    for a in r.action_plan:
        c = t.add_row().cells
        for i, v in enumerate([a.when, a.owner, a.task, a.counterpart, a.message]):
            c[i].text = v
    d.add_heading("[첨부 1] ⑨ PR Stance", level=2)
    d.add_paragraph(f"공식: 「{r.pr_stance.official_message}」")
    d.add_paragraph("강조: " + " / ".join(r.pr_stance.emphasize))
    d.add_paragraph("금기어: " + " · ".join(r.pr_stance.taboo_words))
    d.add_heading("[첨부 2] ⑩ CEO 예상 질문", level=2)
    for i, q in enumerate(r.ceo_qa, 1):
        d.add_paragraph(f"Q{i}. {q.question}").runs[0].bold = True
        d.add_paragraph(q.answer)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
