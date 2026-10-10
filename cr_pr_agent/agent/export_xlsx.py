"""운영 워크북 자동 갱신 — DB(원천) → 엑셀(조회·대시보드용).

시트 구성은 원본 워크북과 같다(01~05) + 06_Dashboard.
04_Decision_Log 의 TOTAL_RISK_SCORE 는 05_Code_Book 가중치를 참조하는 수식, D_DAY 는
DUE_DATE − TODAY() 수식으로 넣는다 (부록 A #10, #11).
"""
from __future__ import annotations

import datetime as dt
import io
from collections import Counter

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import store
from .config import Taxonomy, codebook

HEAD = PatternFill("solid", fgColor="283593")
WHITE = Font(color="FFFFFF", bold=True)
LEVEL_FILL = {"CRITICAL": "F8CBCB", "HIGH": "FCE0C4", "MEDIUM": "FFF3B0", "LOW": "D6EFD8"}

LOG_COLUMNS = store.CORE_COLUMNS + store.EXT_COLUMNS
SCHEMA_DOC = {
    "DOC_ID": ("문서관리번호", "VARCHAR(30)", "DOC-YYMMDD-NNN, 유일키"),
    "UPLOAD_DATE": ("업로드일시", "DATETIME", "문서 수신·업로드 시점"),
    "DOC_TITLE": ("문서제목", "VARCHAR(255)", "원문 문서명 또는 헤드라인"),
    "SOURCE_TYPE": ("출처유형", "ENUM", "정부부처 / 국회 / 경쟁사 / 언론 / 사내보고서"),
    "ISSUE_ID": ("세분류코드", "VARCHAR(10)", "Taxonomy Master 드롭다운"),
    "SCORE_REG": ("규제리스크 점수", "INT 1~5", "규제·인가·제재 영향도 (5=치명)"),
    "SCORE_FIN": ("재무리스크 점수", "INT 1~5", "CapEx·매출·과징금 영향 (5=치명)"),
    "SCORE_PR": ("여론리스크 점수", "INT 1~5", "언론 비판·불매·이미지 타격 (5=치명)"),
    "SCORE_LEG": ("법무리스크 점수", "INT 1~5", "위법성·소송 패소 가능성 (5=치명)"),
    "TOTAL_RISK_SCORE": ("종합리스크 점수", "DECIMAL(4,1)", "(0.35×REG+0.25×PR+0.25×FIN+0.15×LEG)÷5×100"),
    "RISK_LEVEL": ("리스크 등급", "ENUM", "CRITICAL / HIGH / MEDIUM / LOW"),
    "EXECUTIVE_SUMMARY": ("1) 핵심내용", "TEXT", "한줄 판단 + 핵심 사실 3개 이내"),
    "RISK_DESCRIPTION": ("2) Risk/Issue", "TEXT", "진짜 쟁점 + 4대 영역별 위협 요인"),
    "DECISION_POINTS": ("3) 의사결정 포인트", "TEXT", "Option A vs B 및 추천안"),
    "TODO_ACTIONS": ("4) 실행계획", "TEXT", "D+1/3/7 R&R·기한 포함 액션"),
    "PR_STANCE": ("[첨부] PR 스탠스", "TEXT", "공식 메시지 / 강조 / 금기어"),
    "DUE_DATE": ("1차 조치 기한", "DATE", "등급별 기본 기한"),
    "D_DAY": ("마감 D-Day", "INT", "DUE_DATE − 오늘 (수식)"),
    "STATUS": ("진행 상태", "ENUM", "신규접수 / 분석완료 / 대응중 / 종결"),
    "DECISION_TYPE": ("보고 Level", "TEXT", "CEO 보고 / 임원 전결 / 부서 간 조율 / 단순 모니터링"),
    "SCORE_OPP": ("전략적 기회 점수", "INT 1~5", "확장 — 출력 Format ④"),
    "ISSUE_ID_SUB": ("부 코드", "VARCHAR(30)", "확장 — 최대 2개, 쉼표 구분"),
    "CLASS_CONFIDENCE": ("분류 신뢰도", "DECIMAL(3,2)", "확장 — 게이트 1 기준 0.7"),
    "SCENARIOS": ("시나리오", "TEXT", "확장 — A/B/C 효과·부작용"),
    "NEGOTIATION_CARDS": ("협상 카드", "TEXT", "확장 — 출력 Format ⑦"),
    "CEO_QA": ("CEO 예상 질문", "TEXT", "확장 — 6문 30초 답변"),
    "AGENT_RECOMMENDATION": ("에이전트 추천안", "TEXT", "확장 — 조직장 수정 전 원본"),
    "FINAL_DECISION": ("조직장 최종 결정", "TEXT", "확장 — 승인 / 수정 채택 / 기각 + 사유"),
    "REVIEWER": ("검수 담당자", "VARCHAR", "확장"),
    "APPROVER": ("최종 결정자", "VARCHAR", "확장"),
    "PROMPT_VERSION": ("Prompt 버전", "VARCHAR", "확장 — Master·Format·L3"),
    "SOURCE_REF": ("원문 참조", "TEXT", "확장 — 파일명·URL"),
}


def _header(ws, cols, widths=None):
    ws.append(cols)
    for i, _ in enumerate(cols, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEAD, WHITE
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 16)
    ws.freeze_panes = "A2"


def build() -> bytes:
    cb = codebook()
    tx = Taxonomy()
    wb = Workbook()

    # 01 Architecture
    ws = wb.active
    ws.title = "01_Architecture_Flow"
    _header(ws, ["단계", "모듈", "입력", "처리 로직", "출력"], {1: 8, 2: 26, 3: 22, 4: 60, 5: 30})
    for row in [
        ("Step 1", "Document Ingestion", "PDF·HWPX·DOCX·기사", "텍스트·표 추출, 정규화, 개인정보 마스킹, DOC_ID 발번", "구조화된 텍스트"),
        ("Step 2", "Taxonomy Tagging", "추출 텍스트", "키워드 규칙 + 분류 sub-agent 앙상블, 신뢰도 0.7 미만 → 게이트1", "ISSUE_ID·신뢰도"),
        ("Step 3", "Risk Assessment", "분류된 이슈", "리스크 패널 sub-agent 4종 + 기회 1종 병렬, 가중 합산은 시스템 수식", "종합 점수·등급·보고 Level"),
        ("Step 4", "Decision Engine", "패널 결과·과거 사례", "Master+Format+L3 조립 → 10단계 생성 ⟲ 검증·critic 루프", "10단계 판단"),
        ("Step 5", "Executive Delivery", "최종 분석", "원페이지 보고서, Decision Log, 워크북 갱신, 게이트2", "보고서·대시보드"),
    ]:
        ws.append(row)

    # 02 Taxonomy
    ws = wb.create_sheet("02_Taxonomy_Master")
    _header(ws, ["L1", "L2", "L3", "ISSUE_ID", "세분류 이슈명", "전문 Prompt", "상태"], {5: 48, 6: 14})
    for i in tx.data["issues"]:
        p = tx.path(i["issue_id"])
        ws.append([f'{p["l1"]["code"]} {p["l1"]["name"]}', f'{p["l2"]["code"]} {p["l2"]["name"]}',
                   f'{p["l3"]["code"]} {p["l3"]["name"]}', i["issue_id"], i["name"], p["l3"].get("prompt"),
                   i.get("status", "active")])

    # 03 Schema
    ws = wb.create_sheet("03_DB_Schema_Dictionary")
    _header(ws, ["No", "컬럼(영문)", "컬럼(한글)", "타입", "설명"], {1: 5, 2: 24, 3: 18, 4: 14, 5: 60})
    for n, c in enumerate(LOG_COLUMNS, 1):
        ws.append([n, c, *SCHEMA_DOC.get(c, ("", "", ""))])

    # 05 Code Book (04 가 참조하므로 먼저 만든다)
    cbws = wb.create_sheet("05_Code_Book")
    _header(cbws, ["구분", "키", "값", "설명"], {1: 14, 2: 16, 3: 12, 4: 50})
    weight_cells = {}
    for k, v in cb["weights"].items():
        cbws.append(["가중치", k, v, "연 1회 조직장 승인"])
        weight_cells[k] = f"'05_Code_Book'!$C${cbws.max_row}"
    for lv in cb["levels"]:
        cbws.append(["등급", lv["level"], lv["min"], f'{lv["grade"]} · 기본 보고 {lv["report"]} · {lv["principle"]}'])
    for k, v in cb["score_scale"].items():
        cbws.append(["척도(5=치명)", k, int(k), v])
    for s in cb["source_types"]:
        cbws.append(["출처유형", s, "", ""])
    for s in cb["statuses"]:
        cbws.append(["진행상태", s, "", ""])

    # 04 Decision Log
    ws = wb.create_sheet("04_Decision_Log", 3)
    _header(ws, LOG_COLUMNS, {3: 36, 12: 50, 13: 50, 14: 50, 15: 50, 16: 40})
    col = {c: get_column_letter(i) for i, c in enumerate(LOG_COLUMNS, 1)}
    recs = store.list_records(heavy=True)
    for r in recs:
        ws.append([r.get(c) for c in LOG_COLUMNS])
        n = ws.max_row
        if r.get("SCORE_REG") is not None:
            ws[f'{col["TOTAL_RISK_SCORE"]}{n}'] = (
                f'=ROUND(({weight_cells["SCORE_REG"]}*{col["SCORE_REG"]}{n}+{weight_cells["SCORE_PR"]}*{col["SCORE_PR"]}{n}'
                f'+{weight_cells["SCORE_FIN"]}*{col["SCORE_FIN"]}{n}+{weight_cells["SCORE_LEG"]}*{col["SCORE_LEG"]}{n})/5*100,1)')
        if r.get("DUE_DATE"):
            ws[f'{col["DUE_DATE"]}{n}'] = dt.date.fromisoformat(r["DUE_DATE"][:10])
            ws[f'{col["DUE_DATE"]}{n}'].number_format = "yyyy-mm-dd"
            ws[f'{col["D_DAY"]}{n}'] = f'=IF({col["STATUS"]}{n}="종결","",{col["DUE_DATE"]}{n}-TODAY())'
        for c in ("EXECUTIVE_SUMMARY", "RISK_DESCRIPTION", "DECISION_POINTS", "TODO_ACTIONS", "PR_STANCE"):
            ws[f"{col[c]}{n}"].alignment = Alignment(wrap_text=True, vertical="top")
    last = max(ws.max_row, 2)
    rl = f'{col["RISK_LEVEL"]}2:{col["RISK_LEVEL"]}{last}'
    for lv, color in LEVEL_FILL.items():
        ws.conditional_formatting.add(rl, CellIsRule(operator="equal", formula=[f'"{lv}"'],
                                                     fill=PatternFill("solid", fgColor=color)))
    dd = f'{col["D_DAY"]}2:{col["D_DAY"]}{last}'
    ws.conditional_formatting.add(dd, FormulaRule(formula=[f'AND(ISNUMBER({col["D_DAY"]}2),{col["D_DAY"]}2<=3)'],
                                                  fill=PatternFill("solid", fgColor="F8CBCB"), font=Font(bold=True)))
    issue_list = DataValidation(type="list", formula1=f"='02_Taxonomy_Master'!$D$2:$D${len(tx.data['issues'])+1}")
    status_list = DataValidation(type="list", formula1='"' + ",".join(cb["statuses"]) + '"')
    ws.add_data_validation(issue_list)
    ws.add_data_validation(status_list)
    issue_list.add(f'{col["ISSUE_ID"]}2:{col["ISSUE_ID"]}1000')
    status_list.add(f'{col["STATUS"]}2:{col["STATUS"]}1000')

    # 06 Dashboard
    ws = wb.create_sheet("06_Dashboard")
    ws["A1"] = f"CR·PR 의사결정 대시보드 — {dt.date.today()} 기준"
    ws["A1"].font = Font(bold=True, size=14)
    row = 3

    def table(title, counter, keys):
        nonlocal row
        ws.cell(row=row, column=1, value=title).font = Font(bold=True)
        row += 1
        for k in keys:
            ws.cell(row=row, column=1, value=k)
            ws.cell(row=row, column=2, value=counter.get(k, 0))
            if k in LEVEL_FILL:
                ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=LEVEL_FILL[k])
            row += 1
        row += 1

    open_recs = [r for r in recs if r["STATUS"] != "종결"]
    table("진행 중 이슈 — 등급별", Counter(r.get("RISK_LEVEL") for r in open_recs), ["CRITICAL", "HIGH", "MEDIUM", "LOW"])
    table("상태별", Counter(r["STATUS"] for r in recs), cb["statuses"])
    table("조직장 결정", Counter((r.get("FINAL_DECISION") or "미결정").split(":")[0] for r in recs),
          ["승인", "수정 채택", "기각", "미결정"])
    table("L1 분포", Counter(tx.path(r["ISSUE_ID"])["l1"]["name"] for r in recs if r.get("ISSUE_ID") in tx.issues),
          [x["name"] for x in tx.data["l1"]])
    table("월별 유입", Counter((r["UPLOAD_DATE"] or "")[:7] for r in recs),
          sorted({(r["UPLOAD_DATE"] or "")[:7] for r in recs}))
    ws.column_dimensions["A"].width = 28

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
